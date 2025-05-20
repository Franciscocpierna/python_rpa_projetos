from docx import Document
from docx.shared import Pt

from openpyxl import load_workbook
import os
import win32com.client

#Configurando o caminho
nome_arquivo_alunos = "C:\\python_projetos\\python_rpa_projetos\\python_word\\Alunos.xlsx"

#Abrindo o arquivo
planilhaDadosAlunos = load_workbook(nome_arquivo_alunos)

#Selecionando a aba sheet/planilha
sheet_selecionada = planilhaDadosAlunos["Nomes"]




#for = para

for linha in range(2, len(sheet_selecionada["A"]) + 1):

    #Abrindo o arquivo do Word
    #arquivoWord = Document("Certificado1.docx")
    
    word =  win32com.client.Dispatch("word.Application")
    word.Visible = False  # Não mostrar o Word
# abre o documento modelo
    arquivoWord = word.Documents.Open(r"C:\\python_projetos\\python_rpa_projetos\\python_word\\Certificado1.doc")
    

    #Configurando os estilos
    estilo = arquivoWord.styles["Normal"]
    
    #Pegamos o nome do aluno que está na linha corrente
    nomeAluno = sheet_selecionada['A%s' % linha].value

    #for = para
    for paragrafo in arquivoWord.paragraphs:
        #if = se
        # if "@nome" in paragrafo.text:
        #     paragrafo.text = nomeAluno
        #     fonte = estilo.font
        #     fonte.name = "Calibri (Corpo)"
        #     fonte.size = Pt(24)
       
     if "@nome" in paragrafo.Range.Text:
        paragrafo.Range.Text = paragrafo.Range.Text.replace("@nome", nomeAluno)
        # fonte = estilo.font
        # fonte.name = "Calibri (Corpo)"
        # fonte.size = Pt(24)
        # Para alterar a fonte de todo o conteúdo word 2003:
        arquivoWord.Content.Font.Name = "Calibri"
        arquivoWord.Content.Font.Size = 24
        
        

    #caminhoCertificado = "C:\\Users\\55119\\Desktop\\Curso RPA\\Word\\Certificados\\" + nomeAluno + ".docx"
     caminhoCertificado = r"C:\\python_projetos\\python_rpa_projetos\\python_word\\" + nomeAluno + ".doc"
    #Salva o certificado com o nome do aluno
    #arquivoWord.save(caminhoCertificado)
    # Salva como .doc (Word 2003)
     arquivoWord.SaveAs(caminhoCertificado, FileFormat=0)
     arquivoWord.Close()
     word.Quit()

print("Certificados gerados com sucesso!")


'''

import win32com.client

# Caminho do arquivo modelo .doc
caminho_modelo = r"C:\\python_projetos\\python_rpa_projetos\\python_word\\Certificado1.doc"
# Caminho do novo certificado .doc
caminho_saida = r"C:\\python_projetos\\python_rpa_projetos\\python_word\\Amanda Batista Alves.doc"

# Abre o Word
word = win32com.client.Dispatch("Word.Application")
word.Visible = False  # Não mostrar o Word

# Abre o documento modelo
doc = word.Documents.Open(caminho_modelo)
# Apaga todos os parágrafos (linhas) do documento
#doc.Content.Delete()
#doc.PageSetup.TopMargin = 0  # ajuste o valor conforme necessário

# Substitui o marcador @nome pelo nome desejado find não funcionou
# find = doc.Content.Find essas linhas não fez a troca
# find.Text = "@nome"
# find.Replacement.Text = "Amanda Batista Alves"
# find.MatchCase = False  # Ignora maiúsculas/minúsculas
# find.MatchWholeWord = False  # Permite encontrar @nome junto de outros caracteres

#find substituido por for
# Aqui, como exemplo, altera o texto do parágrafo que contém "@nome"
for paragrafo in doc.Paragraphs:
    if "@nome" in paragrafo.Range.Text:
        paragrafo.Range.Text = paragrafo.Range.Text.replace("@nome", "Amanda Batista Alves")


# Aqui, como exemplo, altera toda a fonte do documento:
doc.Content.Font.Name = "Calibri"
doc.Content.Font.Size = 24

# Salva como .doc (Word 2003)
doc.SaveAs(caminho_saida, FileFormat=0)
#essas linhas são para exibir o texto de cada linha/parágrafo
# for paragrafo in doc.Paragraphs:
#     texto = paragrafo.Range.Text
#     print(texto)  # Exibe o texto de cada linha/parágrafo

doc.Close()
word.Quit()

print("Certificado gerado em formato Word 2003 (.doc) com sucesso!")








'''