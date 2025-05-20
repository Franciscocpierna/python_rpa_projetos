from docx import Document
from docx.shared import Pt

from openpyxl import load_workbook
import os

#Configurando o caminho
nome_arquivo_alunos = "C:\\python_projetos\\python_rpa_projetos\\python_word\\Alunos.xlsx"

#Abrindo o arquivo
planilhaDadosAlunos = load_workbook(nome_arquivo_alunos)

#Selecionando a aba sheet/planilha
sheet_selecionada = planilhaDadosAlunos["Nomes"]

#for = para
for linha in range(2, len(sheet_selecionada["A"]) + 1):

    #Abrindo o arquivo do Word
    arquivoWord = Document("C:\\python_projetos\\python_rpa_projetos\\python_word\\Certificado1.docx")

    #Configurando os estilos
    estilo = arquivoWord.styles["Normal"]
    
    #Pegamos o nome do aluno que está na linha corrente
    nomeAluno = sheet_selecionada['A%s' % linha].value

    #for = para
    for paragrafo in arquivoWord.paragraphs:
        #if = se
        if "@nome" in paragrafo.text:
            paragrafo.text = nomeAluno
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

    caminhoCertificado = "C:\\python_projetos\\python_rpa_projetos\\python_word\\" + nomeAluno + ".docx"

    #Salva o certificado com o nome do aluno
    arquivoWord.save(caminhoCertificado)

print("Certificados gerados com sucesso!")